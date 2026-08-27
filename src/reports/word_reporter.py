"""
Word Reporter - Genera reportes en formato Word
"""

from typing import List, Dict
from datetime import datetime
from pathlib import Path
import logging

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

logger = logging.getLogger(__name__)


class WordReporter:
    """Genera reportes en formato .docx"""

    def __init__(self, config: dict):
        """Inicializar reporter"""
        self.config = config
        self.proyecto_nombre = config.get("proyecto", {}).get("nombre", "AjamSi")
        self.isla = config.get("proyecto", {}).get("isla", "Lanzarote")

    def generar_reporte(
        self,
        noticias: List[Dict],
        ideas: List[Dict],
        videos: List[Dict],
        fecha: datetime = None
    ) -> str:
        """
        Generar reporte Word con noticias e ideas

        Args:
            noticias: Lista de noticias scrapeadas
            ideas: Lista de ideas generadas
            videos: Lista de videos compilados
            fecha: Fecha del reporte (default: hoy)

        Returns:
            Ruta del archivo generado
        """
        if not fecha:
            fecha = datetime.now()

        # Crear documento
        doc = Document()

        # Encabezado
        self._agregar_encabezado(doc, fecha)

        # Resumen de ejecución
        self._agregar_resumen(doc, len(noticias), len(ideas), len(videos))

        # Sección de noticias
        self._agregar_noticias(doc, noticias[:20])  # Primeras 20

        # Sección de ideas
        self._agregar_ideas(doc, ideas[:30])  # Primeras 30

        # Guardar
        ruta_output = Path("reportes") / f"ajamsi_reporte_{fecha.strftime('%Y%m%d_%H%M%S')}.docx"
        ruta_output.parent.mkdir(exist_ok=True)

        doc.save(str(ruta_output))
        logger.info(f"✓ Reporte generado: {ruta_output}")

        return str(ruta_output)

    def _agregar_encabezado(self, doc: Document, fecha: datetime):
        """Agregar encabezado al documento"""
        titulo = doc.add_paragraph()
        titulo.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        titulo_run = titulo.add_run(f"📰 {self.proyecto_nombre} - Reporte de Noticias")
        titulo_run.font.size = Pt(24)
        titulo_run.font.bold = True
        titulo_run.font.color.rgb = RGBColor(0, 102, 204)

        # Subtítulo
        subtitulo = doc.add_paragraph()
        subtitulo.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        subtitulo_run = subtitulo.add_run(f"Isla: {self.isla} | {fecha.strftime('%d/%m/%Y %H:%M')}")
        subtitulo_run.font.size = Pt(12)
        subtitulo_run.font.color.rgb = RGBColor(100, 100, 100)

        doc.add_paragraph()  # Espacio

    def _agregar_resumen(self, doc: Document, num_noticias: int, num_ideas: int, num_videos: int):
        """Agregar resumen de ejecución"""
        doc.add_heading("📊 Resumen de Ejecución", level=1)

        tabla = doc.add_table(rows=4, cols=2)
        tabla.style = "Light Grid Accent 1"

        # Headers
        tabla.rows[0].cells[0].text = "Métrica"
        tabla.rows[0].cells[1].text = "Valor"

        # Datos
        tabla.rows[1].cells[0].text = "📰 Noticias Scrapeadas"
        tabla.rows[1].cells[1].text = str(num_noticias)

        tabla.rows[2].cells[0].text = "💡 Ideas Generadas"
        tabla.rows[2].cells[1].text = str(num_ideas)

        tabla.rows[3].cells[0].text = "🎬 Videos Compilados"
        tabla.rows[3].cells[1].text = str(num_videos)

        doc.add_paragraph()  # Espacio

    def _agregar_noticias(self, doc: Document, noticias: List[Dict]):
        """Agregar sección de noticias"""
        doc.add_heading("📰 Noticias Scrapeadas", level=1)

        for i, noticia in enumerate(noticias, 1):
            # Título con número
            titulo = doc.add_heading(f"{i}. {noticia.get('titulo', 'Sin título')}", level=2)

            # Información
            info = doc.add_paragraph()
            info_text = f"Fuente: {noticia.get('fuente', 'N/A')} | "
            info_text += f"Fecha: {noticia.get('fecha_publicacion', 'N/A')}"
            info.add_run(info_text).font.italic = True
            info.runs[0].font.size = Pt(10)

            # Contenido
            contenido = noticia.get('contenido', '')
            if contenido:
                # Limitar a 300 caracteres
                contenido_truncado = contenido[:300] + "..." if len(contenido) > 300 else contenido
                doc.add_paragraph(contenido_truncado)

            # URL
            url = noticia.get('url', '')
            if url:
                url_p = doc.add_paragraph()
                url_run = url_p.add_run(f"🔗 {url}")
                url_run.font.color.rgb = RGBColor(0, 0, 255)
                url_run.font.underline = True
                url_run.font.size = Pt(10)

            doc.add_paragraph()  # Espacio entre noticias

    def _agregar_ideas(self, doc: Document, ideas: List[Dict]):
        """Agregar sección de ideas"""
        doc.add_heading("💡 Ideas Generadas", level=1)

        for i, idea in enumerate(ideas, 1):
            # Título
            titulo = doc.add_heading(f"{i}. {idea.get('titulo', 'Sin título')}", level=2)

            # Metadata
            meta = doc.add_paragraph()
            meta_text = f"Concepto: {idea.get('concepto', 'N/A')} | "
            meta_text += f"Ángulo: {idea.get('angulo_humor', 'N/A')}"
            meta.add_run(meta_text).font.italic = True
            meta.runs[0].font.size = Pt(10)

            # Descripción
            if 'descripcion' in idea:
                doc.add_paragraph(idea['descripcion'])

            # Elementos locales
            elementos = idea.get('elementos_locales', [])
            if elementos:
                elem_p = doc.add_paragraph("Elementos locales: ", style='List Bullet')
                elem_p.add_run(", ".join(elementos))

            # Colores sugeridos
            colores = idea.get('colores', [])
            if colores:
                col_p = doc.add_paragraph("Paleta de colores: ", style='List Bullet')
                col_p.add_run(", ".join(colores))

            doc.add_paragraph()  # Espacio
